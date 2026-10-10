"""Regenerate synthetic, non-personal PDF/DOCX fixtures. No production parser code.

Requires reportlab, python-docx, and pypdf (available in the Codex bundled runtime).
Run this script from any directory. Its output stays beside this script.
"""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader

DEST = Path(__file__).resolve().parent / "fixtures"
DEST.mkdir(exist_ok=True)


def pdf(name, pages):
    output = BytesIO()
    canvas = Canvas(output, invariant=1, pageCompression=1)
    for lines in pages:
        for i, line in enumerate(lines):
            # Keep limit-test text inside the page rather than clipping it horizontally.
            canvas.setFont("Helvetica", 0.1 if len(line) > 500 else 12)
            canvas.drawString(40, 740 - i * 18, line)
        canvas.showPage()
    canvas.save()
    data = output.getvalue()
    (DEST / name).write_bytes(data)
    return data


normal_pdf = pdf("text.pdf", [["Chlorophyll absorbs light energy."], ["Photosynthesis releases oxygen."]])
pdf("partial.pdf", [["Light energy powers photosynthesis."], [], ["Glucose stores chemical energy."]])
pdf("too-many-pages.pdf", [["Synthetic page."]] * 31)
pdf("page-limit.pdf", [["Synthetic page."]] * 30)
pdf("too-much-text.pdf", [["x" * 5000] * 21])
pdf("empty.pdf", [[]])
image_pdf = BytesIO()
canvas = Canvas(image_pdf, invariant=1)
# A synthetic PNG image with no text layer; no personal or scanned material.
from base64 import b64decode
image = b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=")
canvas.drawImage(ImageReader(BytesIO(image)), 40, 600, width=100, height=100)
canvas.showPage()
canvas.save()
(DEST / "image-only.pdf").write_bytes(image_pdf.getvalue())
writer = PdfWriter()
writer.append(PdfReader(BytesIO(normal_pdf)))
writer.encrypt("synthetic-test-password")
with (DEST / "password.pdf").open("wb") as output:
    writer.write(output)
(DEST / "corrupt.pdf").write_bytes(b"%PDF-1.7\nnot a valid PDF\n%%EOF")

doc = Document()
doc.add_paragraph("Photosynthesis requires light energy.")
doc.add_paragraph("Plants do not consume oxygen in this source assertion.")
table = doc.add_table(rows=2, cols=2)
for row, texts in zip(table.rows, [["Input", "Output"], ["Water", "Glucose"]]):
    for cell, text in zip(row.cells, texts):
        cell.text = text
doc.add_paragraph("End of the synthetic lesson.")
doc.save(DEST / "paragraphs-table.docx")
Document().save(DEST / "empty.docx")
long_doc = Document()
long_doc.add_paragraph("x" * 100_001)
long_doc.save(DEST / "too-much-text.docx")
(DEST / "corrupt.docx").write_bytes(b"PK\x03\x04not a valid archive")
with ZipFile(DEST / "not-docx.zip", "w", ZIP_DEFLATED) as archive:
    archive.writestr("readme.txt", "Not a Word document.")
with ZipFile(DEST / "paragraphs-table.docx") as original:
    entries = {entry.filename: original.read(entry.filename) for entry in original.infolist()}
with ZipFile(DEST / "expanded-limit.docx", "w", ZIP_DEFLATED) as archive:
    for name, data in entries.items():
        archive.writestr(name, data)
    archive.writestr("word/media/padding.bin", b"0" * 10_000_001)
with ZipFile(DEST / "malformed-xml.docx", "w", ZIP_DEFLATED) as archive:
    for name, data in entries.items():
        archive.writestr(name, b"<not-valid" if name == "word/document.xml" else data)

print(f"Generated {len(list(DEST.iterdir()))} synthetic fixtures in {DEST}")
